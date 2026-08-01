from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


SEARCH_TOKEN = "MATECLAW_FIXTURE_ALPHA_2026"


def set_style_font(style, name: str, size: float, color: str | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def build_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, "Calibri", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        set_style_font(style, "Calibri", size, color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "MateClaw P0 Integration Fixture"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_style_font(header.style, "Calibri", 9, "6B7280")

    footer = section.footer.paragraphs[0]
    footer.text = "Deterministic development fixture - safe to reset"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_style_font(footer.style, "Calibri", 9, "6B7280")

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("MateClaw Cloud Integration Fixture")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subrun = subtitle.add_run("DOCX sample for MCP read, search, summarize, and preview checks")
    subrun.italic = True
    subrun.font.color.rgb = RGBColor.from_string("4B5563")

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This small deterministic document belongs to the isolated P0 fixture folder. "
        "It contains stable text for repeatable cloud-drive and MCP integration checks."
    )

    doc.add_heading("Searchable content", level=2)
    doc.add_paragraph(
        f"Stable token: {SEARCH_TOKEN}. The integration should preserve tenant identity, "
        "return structured errors, and expose a correlation trace ID."
    )

    doc.add_heading("Expected checks", level=2)
    doc.add_paragraph(
        "The file can be listed, searched, previewed, summarized, versioned, and moved "
        "without relying on production business documents."
    )

    doc.core_properties.title = "MateClaw Cloud Integration Fixture"
    doc.core_properties.subject = "P0 deterministic MCP fixture"
    doc.core_properties.author = "MateClaw Integration Test"
    doc.save(path)

    reopened = Document(path)
    text = "\n".join(p.text for p in reopened.paragraphs)
    if SEARCH_TOKEN not in text or len(reopened.sections) != 1:
        raise RuntimeError("DOCX structural verification failed")


def build_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    title = styles["Title"]
    title.fontName = "Helvetica-Bold"
    title.fontSize = 22
    title.textColor = "#0B2545"
    title.spaceAfter = 18
    heading = styles["Heading2"]
    heading.fontName = "Helvetica-Bold"
    heading.fontSize = 13
    heading.textColor = "#2E74B5"
    heading.spaceBefore = 12
    heading.spaceAfter = 7
    body = styles["BodyText"]
    body.fontName = "Helvetica"
    body.fontSize = 11
    body.leading = 15
    body.spaceAfter = 8

    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="MateClaw Cloud Integration Fixture",
        author="MateClaw Integration Test",
    )
    story = [
        Paragraph("MateClaw Cloud Integration Fixture", title),
        Paragraph("Purpose", heading),
        Paragraph(
            "Deterministic PDF sample for cloud-drive preview and MCP read checks.",
            body,
        ),
        Paragraph("Searchable content", heading),
        Paragraph(
            f"Stable token: {SEARCH_TOKEN}. This file contains no production data.",
            body,
        ),
        Spacer(1, 12),
        Paragraph(
            "Expected result: tenant-scoped access, structured errors, and traceable tool calls.",
            body,
        ),
    ]
    document.build(story)

    reader = PdfReader(path)
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    if len(reader.pages) != 1 or SEARCH_TOKEN not in extracted:
        raise RuntimeError("PDF structural verification failed")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output_dir")
    args = parser.parse_args()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    build_docx(output_dir / "fixture-sample.docx")
    build_pdf(output_dir / "fixture-sample.pdf")
    print(output_dir)


if __name__ == "__main__":
    main()
